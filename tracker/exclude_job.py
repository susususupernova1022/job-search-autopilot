#!/usr/bin/env python3
"""exclude_job.py — record jobs you reviewed and rejected, so the scan never shows them again.

Key = company | title | job ID (normalized). A repost with a new ID will show up once more — deliberate:
better to see a job twice than to silently lose one.

Usage:
  python3 tracker/exclude_job.py --company "Mercury" --title "Product Manager, Payments" --job-id 4512345 \
      --url "https://..." --reason "JD requires 8+ years"
  echo '[{"company":"Mercury","title":"PM, Payments","job_id":"451","reason":"too senior"}]' | python3 tracker/exclude_job.py --stdin

Writes only the Excluded sheet. Existing keys are skipped and reported — no need to check first.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date

import openpyxl

from common import EXCLUDED_HEADERS, EXCLUDED_SHEET, find_tracker, first_blank_row, make_key, read_rows


def ensure_sheet(wb):
    if EXCLUDED_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(EXCLUDED_SHEET)
        ws.append(EXCLUDED_HEADERS)
        return ws
    return wb[EXCLUDED_SHEET]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--company")
    ap.add_argument("--title")
    ap.add_argument("--job-id", default="")
    ap.add_argument("--url", default="")
    ap.add_argument("--reason", default="Excluded after review")
    ap.add_argument("--stdin", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if args.stdin:
        items = json.load(sys.stdin)
    elif args.company and args.title:
        items = [{"company": args.company, "title": args.title, "job_id": args.job_id, "url": args.url, "reason": args.reason}]
    else:
        sys.exit("Need --company and --title, or --stdin with a JSON list")

    tracker = find_tracker()
    wb = openpyxl.load_workbook(tracker)
    ws = ensure_sheet(wb)
    existing = {make_key(r.get("Company"), r.get("Title"), r.get("Job ID")) for r in read_rows(wb, EXCLUDED_SHEET)}

    added, skipped = [], []
    row = first_blank_row(ws, len(EXCLUDED_HEADERS))
    for item in items:
        company = (item.get("company") or "").strip()
        title = (item.get("title") or "").strip()
        job_id = str(item.get("job_id") or "").strip()
        if not company or not title:
            skipped.append({"item": item, "why": "missing company or title"})
            continue
        key = make_key(company, title, job_id)
        if key in existing:
            skipped.append({"item": f"{company} / {title}", "why": "already excluded"})
            continue
        values = [item.get("date") or date.today().isoformat(), company, title, job_id,
                  (item.get("url") or "").strip(), (item.get("reason") or "Excluded after review").strip()]
        if not args.dry_run:
            for col, v in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=v)
        existing.add(key)
        added.append(f"{company} / {title}" + (f" ({job_id})" if job_id else ""))
        row += 1

    if args.dry_run:
        print(json.dumps({"would_add": added, "skipped": skipped}, ensure_ascii=False, indent=2))
        return
    wb.save(tracker)
    print(json.dumps({"tracker": str(tracker), "added": added, "skipped": skipped}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
