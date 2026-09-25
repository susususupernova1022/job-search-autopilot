#!/usr/bin/env python3
"""mark_in_progress.py — reserve picked jobs in the Tracker as "In progress" the moment you pick them.

Why: applying happens later, one new session per lane, sometimes days later. Writing the row now means
tomorrow's scan dedupes it and doesn't show you the same job again.

Writes columns A–G only (Date, Company, Title, Job ID, Method, Status, Lane). Never touches H onward.
Status = "In progress | <lane> | <YYYY-MM-DD>". Date = the day you picked it; clear_in_progress.py
changes it to the real submit date.

Usage:
  python3 tracker/mark_in_progress.py --company "Plaid" --title "Product Manager, Credit" --job-id 123 --lane PM --method Ashby
  echo '[{"company":"Plaid","title":"PM, Credit","job_id":"123","lane":"PM","method":"Ashby"}]' | python3 tracker/mark_in_progress.py --stdin
  add --dry-run to preview.

Use the company name as you write it in the Tracker, not the board slug.
Rows whose Company|Title already exist in Applications are skipped (no double reservation).
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date

import openpyxl

from common import (
    APP_HEADERS, APPLICATIONS_SHEET, IN_PROGRESS, SEP, find_tracker, first_blank_row, load_config,
    make_loose_key, read_rows,
)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--company")
    ap.add_argument("--title")
    ap.add_argument("--job-id", default="")
    ap.add_argument("--lane", default="")
    ap.add_argument("--method", default="TBD", help="ATS / channel")
    ap.add_argument("--date", default="", help="YYYY-MM-DD, default today")
    ap.add_argument("--stdin", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if args.stdin:
        items = json.load(sys.stdin)
    elif args.company and args.title:
        items = [{"company": args.company, "title": args.title, "job_id": args.job_id, "lane": args.lane, "method": args.method}]
    else:
        sys.exit("Need --company and --title, or --stdin with a JSON list")

    lanes = load_config()["lanes"]
    today = args.date.strip() or date.today().isoformat()
    tracker = find_tracker()
    wb = openpyxl.load_workbook(tracker)
    if APPLICATIONS_SHEET not in wb.sheetnames:
        sys.exit(f"Tracker has no '{APPLICATIONS_SHEET}' sheet — wrong file?")
    ws = wb[APPLICATIONS_SHEET]
    existing = {make_loose_key(r.get("Company"), r.get("Title")) for r in read_rows(wb, APPLICATIONS_SHEET)}

    added, skipped = [], []
    row = first_blank_row(ws, len(APP_HEADERS))
    for item in items:
        company = (item.get("company") or "").strip()
        title = (item.get("title") or "").strip()
        lane = (item.get("lane") or "").strip()
        if not company or not title:
            skipped.append({"item": item, "why": "missing company or title"})
            continue
        if lane and lane not in lanes:
            skipped.append({"item": f"{company} / {title}", "why": f"lane '{lane}' is not one of {lanes}"})
            continue
        key = make_loose_key(company, title)
        if key in existing:
            skipped.append({"item": f"{company} / {title}", "why": "already in Applications"})
            continue
        picked = (item.get("date") or today).strip()
        status = SEP.join([IN_PROGRESS, lane or "no lane", picked])
        values = [picked, company, title, str(item.get("job_id") or "").strip(), (item.get("method") or "TBD").strip(), status, lane]
        if not args.dry_run:
            for col, value in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=value)
        existing.add(key)
        added.append({"row": row, "company": company, "title": title, "lane": lane, "status": status})
        row += 1

    if args.dry_run:
        print(json.dumps({"would_add": added, "skipped": skipped}, ensure_ascii=False, indent=2))
        return
    wb.save(tracker)
    print(json.dumps({"tracker": str(tracker), "added": added, "skipped": skipped}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
