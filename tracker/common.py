"""Shared helpers for the tracker scripts: config, file lookup, key normalization, sheet reading.

All scripts locate files relative to the repo root (the folder containing autopilot_config.json),
so they work the same on your Mac and inside a Claude session that has the folder connected.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl  # noqa: F401
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl  (add --break-system-packages inside a Claude session)")

REPO_ROOT = Path(__file__).resolve().parents[1]
CONFIG_FILE = REPO_ROOT / "autopilot_config.json"

DEFAULT_CONFIG = {
    "tracker_file": "data/Application_Tracker.xlsx",
    "question_bank_file": "data/Question_Bank.xlsx",
    "bullet_db_file": "data/Bullet_Database.xlsx",
    "target_companies_file": "data/Target_Companies.xlsx",
    "lanes": ["Finance/FP&A", "Strategy&Ops", "PM", "PMM/Marketing"],
    "in_progress_stale_days": 3,
}

# ---- Tracker layout -------------------------------------------------------
APPLICATIONS_SHEET = "Applications"
EXCLUDED_SHEET = "Excluded"
QUOTA_SHEET = "Company_Quota"

# The pipeline owns columns A–G. Columns H onward are yours (outcome, interview stage, your own notes …)
# and no script ever writes to them.
APP_HEADERS = ["Date", "Company", "Title", "Job ID", "Method", "Status", "Lane"]
USER_HEADERS = ["Outcome", "Fit %", "Notes"]
EXCLUDED_HEADERS = ["Date", "Company", "Title", "Job ID", "URL", "Reason"]
QUOTA_HEADERS = ["Company", "Aliases", "Max applications", "Window days", "Active", "Notes"]

# ---- Status vocabulary (prefix of the Status column) ------------------------
#   In progress | <lane> | <date>  → picked, being applied in a lane session. Counts as applied (blocks
#                                    re-scan, uses quota). Reported as stale after N days.
#   Submitted …                    → applied.
#   Not applied | <reason>         → abandoned. Does NOT count as applied; the job can resurface in scans.
#   Skipped | <reason>             → permanent no (e.g. JD says no sponsorship). Counts as applied so it
#                                    never resurfaces. Use this, not "Not applied", for hard knockouts.
IN_PROGRESS = "In progress"
SUBMITTED = "Submitted"
NOT_APPLIED = "Not applied"
SKIPPED = "Skipped"
SEP = " | "


def load_config() -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_FILE.exists():
        cfg.update(json.loads(CONFIG_FILE.read_text(encoding="utf-8")))
    return cfg


def data_path(key: str) -> Path:
    return REPO_ROOT / load_config()[key]


def find_tracker() -> Path:
    path = data_path("tracker_file")
    if not path.exists():
        sys.exit(f"Tracker not found at {path}. Run: python3 tracker/init_workbooks.py  (or check the folder is connected)")
    return path


_WS = re.compile(r"\s+")
_PUNCT = re.compile(r"[^\w\s&+/-]", re.UNICODE)
_LEGAL_SUFFIX = re.compile(
    r"\b(inc|llc|ltd|limited|corp|corporation|co|company|holdings|technologies|technology|labs|group|plc|gmbh|pte|sa|nv|ag)\b",
    re.IGNORECASE,
)


def norm(value) -> str:
    """Must match norm() in scanner/jd_scanner.js."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().lower()
    text = re.sub(r"\([^)]*\)", " ", text)
    text = _PUNCT.sub(" ", text)
    text = _LEGAL_SUFFIX.sub(" ", text)
    return _WS.sub(" ", text).strip()


# Job IDs get written as "Greenhouse 7818068003" or "req R0007198"; the scanner sees the bare ID.
# Strip the source prefix or Excluded matching never hits.
_ID_SOURCE_PREFIX = re.compile(
    r"^(greenhouse|gh|lever|ashby|ashbyhq|workday|myworkday|indeed|linkedin|icims|smartrecruiters|jobvite|"
    r"phenom|taleo|adp|bamboohr|successfactors|careers?|job|req|requisition|id)[\s:_-]+"
)


def norm_job_id(value) -> str:
    text = norm(value)
    if not text or text in {"none", "n a", "na", "-"}:
        return "-"
    prev = None
    while prev != text:
        prev = text
        text = _ID_SOURCE_PREFIX.sub("", text).strip()
    return text or "-"


def make_key(company, title, job_id) -> str:
    return f"{norm(company)}|{norm(title)}|{norm_job_id(job_id)}"


def make_loose_key(company, title) -> str:
    return f"{norm(company)}|{norm(title)}"


def parse_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_rows(wb, sheet_name: str) -> list[dict]:
    if sheet_name not in wb.sheetnames:
        return []
    rows = list(wb[sheet_name].iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        out.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    return out


def first_blank_row(ws, width: int) -> int:
    """First truly empty row (openpyxl's max_row counts formatted-but-empty rows)."""
    row = ws.max_row
    while row > 1:
        if any(ws.cell(row=row, column=c).value not in (None, "") for c in range(1, width + 1)):
            return row + 1
        row -= 1
    return 2


def status_of(row: dict) -> str:
    return str(row.get("Status") or "").strip()


def split_status(status: str) -> list[str]:
    return [p.strip() for p in re.split(r"[|｜]", status or "") if p.strip()]
