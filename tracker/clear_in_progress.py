#!/usr/bin/env python3
"""clear_in_progress.py — close out an "In progress" row at the end of a lane session.

  (default)            Status → "Submitted | <lane>", Date → submit date (today unless --date)
  --abandon "<why>"    Status → "Not applied | <why>". The job can resurface in future scans.
  --skip "<why>"       Status → "Skipped | <why>". Permanent: never resurfaces (use for hard knockouts).
  --status "<text>"    write any status you like

Touches only column A (Date) and column F (Status). Never H onward.

Usage:
  python3 tracker/clear_in_progress.py --company "Plaid" --title "Product Manager, Credit"
  python3 tracker/clear_in_progress.py --company "Plaid" --title "..." --abandon "portal needs native upload; retry manually"
  echo '[{"company":"Plaid","title":"PM, Credit"}]' | python3 tracker/clear_in_progress.py --stdin

Tip: for Workday, take the submit date from the candidate home page ("My Applications"), not the
day you ran this — sessions sometimes close out a day late.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date

import openpyxl

from common import (
    APPLICATIONS_SHEET, IN_PROGRESS, NOT_APPLIED, SEP, SKIPPED, SUBMITTED, find_tracker, make_loose_key,
    norm_job_id, split_status,
)

COL_DATE, COL_COMPANY, COL_TITLE, COL_JOB_ID, COL_STATUS, COL_LANE = 1, 2, 3, 4, 6, 7


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--company")
    ap.add_argument("--title")
    ap.add_argument("--job-id", default="", help="disambiguate when several rows match")
    ap.add_argument("--date", default="", help="actual submit date YYYY-MM-DD, default today")
    ap.add_argument("--abandon", default="")
    ap.add_argument("--skip", default="")
    ap.add_argument("--status", default="")
    ap.add_argument("--stdin", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if args.stdin:
        items = json.load(sys.stdin)
    elif args.company and args.title:
        items = [{"company": args.company, "title": args.title, "job_id": args.job_id, "date": args.date,
                  "abandon": args.abandon, "skip": args.skip, "status": args.status}]
    else:
        sys.exit("Need --company and --title, or --stdin with a JSON list")

    today = args.date.strip() or date.today().isoformat()
    tracker = find_tracker()
    wb = openpyxl.load_workbook(tracker)
    ws = wb[APPLICATIONS_SHEET]

    updated, not_found = [], []
    for item in items:
        company = (item.get("company") or "").strip()
        title = (item.get("title") or "").strip()
        if not company or not title:
            not_found.append({"item": item, "why": "missing company or title"})
            continue
        want_key = make_loose_key(company, title)
        want_id = norm_job_id(item.get("job_id") or "")
        hits = []
        for r in range(2, ws.max_row + 1):
            status = str(ws.cell(row=r, column=COL_STATUS).value or "").strip()
            if not status.startswith(IN_PROGRESS):
                continue
            if make_loose_key(ws.cell(row=r, column=COL_COMPANY).value, ws.cell(row=r, column=COL_TITLE).value) != want_key:
                continue
            if want_id != "-" and norm_job_id(ws.cell(row=r, column=COL_JOB_ID).value) != want_id:
                continue
            hits.append(r)
        if not hits:
            not_found.append({"item": f"{company} / {title}", "why": "no 'In progress' row found"})
            continue

        for r in hits:
            old = str(ws.cell(row=r, column=COL_STATUS).value or "")
            parts = split_status(old)
            lane = ws.cell(row=r, column=COL_LANE).value or (parts[1] if len(parts) > 1 else "")
            abandon, skip, custom = (item.get("abandon") or "").strip(), (item.get("skip") or "").strip(), (item.get("status") or "").strip()
            if custom:
                new = custom
            elif abandon:
                new = SEP.join([NOT_APPLIED, abandon])
            elif skip:
                new = SEP.join([SKIPPED, skip])
            else:
                new = SEP.join([p for p in (SUBMITTED, lane) if p])
            submitted = not (abandon or skip or custom)
            when = (item.get("date") or today) if submitted else None
            if not args.dry_run:
                ws.cell(row=r, column=COL_STATUS, value=new)
                if when:
                    ws.cell(row=r, column=COL_DATE, value=when)
            updated.append({"row": r, "company": company, "title": title, "old": old, "new": new, "date": when})

    if args.dry_run:
        print(json.dumps({"would_update": updated, "not_found": not_found}, ensure_ascii=False, indent=2))
        return
    wb.save(tracker)
    print(json.dumps({"tracker": str(tracker), "updated": updated, "not_found": not_found}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
