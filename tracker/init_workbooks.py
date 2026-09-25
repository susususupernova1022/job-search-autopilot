#!/usr/bin/env python3
"""init_workbooks.py — create the empty data workbooks with the right sheets and headers.

Creates (never overwrites an existing file unless --force):
  data/Application_Tracker.xlsx   Applications / Excluded / Company_Quota
  data/Question_Bank.xlsx         Application Question Bank (seeded) / Work & Education Timeline
  data/Bullet_Database.xlsx       Bullets
  data/Target_Companies.xlsx      Target Companies / Removed
and the output folders outputs/CV and outputs/Cover Letter.

Options:
  --with-sample-companies   pre-fill Target_Companies from templates/target_companies_sample.csv
  --with-sample-quotas      pre-fill Company_Quota from templates/company_quota_sample.csv
  --force                   overwrite existing workbooks (careful)
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from common import (
    APP_HEADERS, APPLICATIONS_SHEET, EXCLUDED_HEADERS, EXCLUDED_SHEET, QUOTA_HEADERS, QUOTA_SHEET,
    REPO_ROOT, USER_HEADERS, data_path,
)

TEMPLATES = REPO_ROOT / "templates"
HEAD_FONT = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE7F3")
USER_FILL = PatternFill("solid", fgColor="EFEFEF")


def style_header(ws, widths=None, user_from=None):
    for i, cell in enumerate(ws[1], start=1):
        cell.font = HEAD_FONT
        cell.fill = USER_FILL if (user_from and i >= user_from) else HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def read_csv(name):
    path = TEMPLATES / name
    if not path.exists():
        return [], []
    with path.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def save(wb, path: Path, force: bool, made: list):
    if path.exists() and not force:
        made.append(f"kept existing {path.relative_to(REPO_ROOT)}")
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    made.append(f"created {path.relative_to(REPO_ROOT)}")


def tracker(sample_quotas: bool):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = APPLICATIONS_SHEET
    ws.append(APP_HEADERS + USER_HEADERS)
    style_header(ws, [12, 22, 40, 18, 16, 36, 16, 16, 8, 40], user_from=len(APP_HEADERS) + 1)
    ws.cell(row=1, column=len(APP_HEADERS) + len(USER_HEADERS) + 1,
            value="← grey columns are yours; scripts only write A–G")
    ex = wb.create_sheet(EXCLUDED_SHEET)
    ex.append(EXCLUDED_HEADERS)
    style_header(ex, [12, 22, 40, 18, 50, 50])
    q = wb.create_sheet(QUOTA_SHEET)
    q.append(QUOTA_HEADERS)
    style_header(q, [22, 24, 16, 12, 8, 50])
    if sample_quotas:
        _, rows = read_csv("company_quota_sample.csv")
        for r in rows:
            q.append([r[0], r[1], int(r[2]), int(r[3]), r[4], r[5]])
    return wb


def question_bank():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Application Question Bank"
    header, rows = read_csv("question_bank_seed.csv")
    ws.append(header or ["Category", "Question (common phrasing)", "Type", "What they screen for", "Rule / how to answer", "My answer", "Notes"])
    for r in rows:
        ws.append(r)
    style_header(ws, [16, 60, 12, 30, 60, 40, 30])
    ws.cell(row=1, column=9, value="Type: fact = answer from profile/timeline automatically | judgment = ask once, then reuse | per-company = ask every time | consent = you tick it yourself")
    tl = wb.create_sheet("Work & Education Timeline")
    tl.append(["Type", "Organization (legal name)", "Title / Degree", "Field of study", "Start (MM/YYYY)", "End (MM/YYYY or Present)",
               "City", "Country", "Paid?", "Sub-role of", "Notes"])
    style_header(tl, [10, 32, 32, 24, 12, 16, 16, 14, 8, 20, 40])
    return wb


def bullet_db():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bullets"
    ws.append(["ID", "Company", "Role", "Resume Version", "Skill Category", "Lane", "Metric", "Bullet Text", "Notes"])
    style_header(ws, [16, 24, 24, 18, 18, 22, 18, 90, 40])
    return wb


def target_companies(sample: bool):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Target Companies"
    header, rows = read_csv("target_companies_sample.csv")
    ws.append(header)
    if sample:
        for r in rows:
            ws.append(r)
    style_header(ws, [22, 16, 16, 14, 32, 22, 11, 6, 20, 24, 10, 60])
    rm = wb.create_sheet("Removed")
    rm.append(["Company", "Reason removed", "Details", "Date"])
    style_header(rm, [22, 24, 60, 12])
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--with-sample-companies", action="store_true")
    ap.add_argument("--with-sample-quotas", action="store_true")
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    made = []
    save(tracker(args.with_sample_quotas), data_path("tracker_file"), args.force, made)
    save(question_bank(), data_path("question_bank_file"), args.force, made)
    save(bullet_db(), data_path("bullet_db_file"), args.force, made)
    save(target_companies(args.with_sample_companies), data_path("target_companies_file"), args.force, made)
    for sub in ("outputs/CV", "outputs/Cover Letter", "data/resume_masters", "data/cover_letter_masters"):
        (REPO_ROOT / sub).mkdir(parents=True, exist_ok=True)
    print("\n".join(made))


if __name__ == "__main__":
    main()
